import streamlit as st
import pandas as pd
from openpyxl import load_workbook
import statistics
import io

st.set_page_config(page_title="Vandføring fra Hipdata.dk", layout="centered")

st.title("Vandføring fra Hipdata.dk")
st.write("Upload CSV fra hipdata.dk og få konverteret Excel med beregninger.")

uploaded_file = st.file_uploader("Vælg CSV-fil", type="csv")

if uploaded_file:
    try:
        df = pd.read_csv(uploaded_file, sep=';', encoding='utf-8-sig')

        if df.empty:
            st.error("CSV-filen er tom.")
        else:
            # Konverter kolonne B til tal
            df.iloc[:, 1] = pd.to_numeric(df.iloc[:, 1].astype(str).str.replace(',', '.'), errors='coerce')

            # Indsæt MONTH-kolonne
            df['Month'] = pd.to_datetime(df.iloc[:, 0], dayfirst=True, errors='coerce').dt.month

            # Beregn middel for maj-sept
            maj_sept_values = df.loc[df['Month'].between(5, 9), df.columns[1]]
            if not maj_sept_values.empty:
                average = round(maj_sept_values.mean() * 1000, 6)
            else:
                average = "Ingen data"

            # Beregn median af maksimum for vintermåneder
            vintermåneder = [1, 2, 3, 4, 10, 11, 12]
            max_values = []
            for m in vintermåneder:
                m_values = df.loc[df['Month'] == m, df.columns[1]]
                if not m_values.empty:
                    max_values.append(m_values.max())

            if max_values:
                median_max = round(statistics.median(max_values) * 1000, 6)
            else:
                median_max = "Ingen data"

            # Opret Excel-fil i hukommelsen
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Sheet1')
                workbook = writer.book
                worksheet = writer.sheets['Sheet1']

                # Skriv middel og median i Excel
                worksheet.cell(row=1, column=5).value = 'Middel (maj–sept)'
                worksheet.cell(row=2, column=5).value = average if isinstance(average, (int, float)) else average

                worksheet.cell(row=1, column=6).value = 'Median af max (vintermåneder)'
                worksheet.cell(row=2, column=6).value = median_max if isinstance(median_max, (int, float)) else median_max

            output.seek(0)

            # Vis resultater i Streamlit
            st.success(f"Middel (maj–sept): {average} L/s" if isinstance(average, (int, float)) else "Ingen data til beregning")
            st.success(f"Median maksimum (vintermåneder): {median_max} L/s" if isinstance(median_max, (int, float)) else "Ingen data til median af max")

            st.download_button(
                label="Download Excel-fil",
                data=output,
                file_name="konverteret.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

    except Exception as e:
        st.error(f"Noget gik galt: {e}")

